import csv
from pathlib import Path
import zipfile

from openpyxl import load_workbook
import pytest


def test_roundtrip_literals_and_no_formulas(tmp_path, csv_tool):
    source, target = tmp_path / "数据.csv", tmp_path / "result.xlsx"
    values = ["00123", "12345678901234567890", "=1+2", "+123", "@name", "#N/A", "1E10",
              "2026-10-22", "https://example.org", "中文🙂", "line1\nline2", 'a,"b"', "  spaces  "]
    with source.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow([f"c{i}" for i in range(len(values))])
        writer.writerow(values)
    original = source.read_bytes()
    assert csv_tool.convert(source, target)["records"] == 2
    workbook = load_workbook(target)
    assert [cell.value for cell in workbook.active[2]] == values
    assert all(cell.data_type == "s" for cell in workbook.active[2])
    assert all(cell.number_format == "@" and cell.hyperlink is None for cell in workbook.active[2])
    assert workbook.active.freeze_panes == "A2"
    workbook.close()
    with zipfile.ZipFile(target) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml")
        assert b"<f>" not in xml
        assert not any("externalLink" in name for name in archive.namelist())
    assert source.read_bytes() == original


@pytest.mark.parametrize("text", ["", "a,b\n1\n", 'a,b\n"unterminated,2', "a\n\x01\n", "a\n" + "x" * 32768,
                                      "a\n" + "🙂" * 16384, "a\n\uffff\n"],
                         ids=["empty", "ragged", "unclosed-quote", "control", "long-text", "long-utf16", "invalid-xml"])
def test_bad_input_leaves_no_output(tmp_path, csv_tool, text):
    source, target = tmp_path / "in.csv", tmp_path / "out.xlsx"
    source.write_text(text, encoding="utf-8")
    with pytest.raises((ValueError, csv.Error)):
        csv_tool.convert(source, target)
    assert not target.exists()
    assert list(tmp_path.iterdir()) == [source]


def test_refuses_overwrite_and_input_path(tmp_path, csv_tool):
    source = tmp_path / "data.csv"
    source.write_text("a\n01\n")
    target = tmp_path / "out.xlsx"
    target.write_bytes(b"keep")
    with pytest.raises(FileExistsError):
        csv_tool.convert(source, target)
    with pytest.raises(ValueError):
        csv_tool.convert(source, source)
    assert target.read_bytes() == b"keep"


def test_explicit_encoding_delimiter_and_no_header(tmp_path, csv_tool):
    source, target = tmp_path / "gbk.tsv", tmp_path / "out.xlsx"
    source.write_bytes("编号\t名称\n0001\t中文\n".encode("gb18030"))
    with pytest.raises(UnicodeError):
        csv_tool.convert(source, target)
    csv_tool.convert(source, target, encoding="gb18030", delimiter="\t", header=False)
    workbook = load_workbook(target)
    assert workbook.active["A2"].value == "0001"
    assert workbook.active["B2"].value == "中文"
    assert workbook.active.freeze_panes is None
    workbook.close()


def test_empty_cells_keep_column_positions(tmp_path, csv_tool):
    source, target = tmp_path / "in.csv", tmp_path / "out.xlsx"
    source.write_text("a,b,c\n,001,\n", encoding="utf-8")
    csv_tool.convert(source, target)
    workbook = load_workbook(target)
    assert workbook.active.max_column == 3
    assert workbook.active["B2"].value == "001"
    assert workbook.active["A2"].value is None
    workbook.close()
