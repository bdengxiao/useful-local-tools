"""Convert delimited text to XLSX with every cell stored as literal text."""
from __future__ import annotations

import argparse
import csv
import json
import os
from pathlib import Path
import re
import shutil
import sys
import tempfile


def convert(source, destination, *, encoding="utf-8-sig", delimiter=",", header=True):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    source, destination = Path(source), Path(destination)
    if source.resolve() == destination.resolve():
        raise ValueError("Input and output must be different files.")
    if destination.suffix.lower() != ".xlsx":
        raise ValueError("Output must have an .xlsx extension.")
    if destination.exists():
        raise FileExistsError("Output exists; choose a new filename.")
    if len(delimiter) != 1 or delimiter in '\r\n"':
        raise ValueError("Delimiter must be one character other than a newline or quote.")
    # Encoding errors are fatal: never replace unknown bytes and corrupt an ID.
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Data")
    if header:
        sheet.freeze_panes = "A2"
    rows, width = 0, None
    fd, temporary = tempfile.mkstemp(suffix=".xlsx", dir=destination.parent)
    os.close(fd)
    saved = False
    old_limit = csv.field_size_limit()
    csv.field_size_limit(1_000_000)
    try:
        with source.open("r", encoding=encoding, newline="") as stream:
            reader = csv.reader(stream, delimiter=delimiter, strict=True)
            for row in reader:
                rows += 1
                if rows > 1_048_576:
                    raise ValueError("Input exceeds Excel's row limit.")
                if width is None:
                    width = len(row)
                    if not width or width > 16_384:
                        raise ValueError("Invalid column count for Excel.")
                if len(row) != width:
                    raise ValueError(f"Record {rows} has {len(row)} columns; expected {width}.")
                cells = []
                for col, value in enumerate(row, 1):
                    if len(value.encode("utf-16-le")) // 2 > 32_767:
                        raise ValueError(f"Record {rows}, column {col} exceeds Excel's text limit.")
                    if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]", value):
                        raise ValueError(f"Record {rows}, column {col} contains unsupported control characters.")
                    cell = WriteOnlyCell(sheet, value=value)
                    cell.data_type = "s"  # Even '=...', '#N/A', and URLs remain literal strings.
                    cell.number_format = "@"
                    if header and rows == 1:
                        cell.font = Font(bold=True, color="FFFFFF")
                        cell.fill = PatternFill("solid", fgColor="174E52")
                    cells.append(cell)
                sheet.append(cells)
        if rows == 0:
            raise ValueError("Input is empty.")
        if header:
            sheet.auto_filter.ref = f"A1:{get_column_letter(width)}{rows}"
        workbook.save(temporary)
        saved = True
        # Exclusive creation refuses overwrites, including a file created during conversion.
        created = False
        try:
            with destination.open("xb") as target, open(temporary, "rb") as result:
                created = True
                shutil.copyfileobj(result, target)
        except BaseException:
            if created:
                destination.unlink(missing_ok=True)
            raise
        return {"ok": True, "records": rows, "columns": width, "all_cells": "text"}
    finally:
        csv.field_size_limit(old_limit)
        if not saved:
            # Finalize streaming worksheet resources even after malformed input.
            try:
                workbook.save(temporary)
            except Exception:
                pass
        workbook.close()
        Path(temporary).unlink(missing_ok=True)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--encoding", default="utf-8-sig", help="Explicit input encoding, e.g. gb18030")
    parser.add_argument("--delimiter", default=",", help="One character, or 'tab'")
    parser.add_argument("--no-header", action="store_true")
    args = parser.parse_args(argv)
    try:
        report = convert(args.input, args.output, encoding=args.encoding,
                         delimiter="\t" if args.delimiter == "tab" else args.delimiter,
                         header=not args.no_header)
    except ImportError:
        print("Install openpyxl: python -m pip install 'openpyxl>=3.1.5,<4'", file=sys.stderr)
        return 2
    except UnicodeError:
        print("Cannot decode input. Supply the correct --encoding; no output was written.", file=sys.stderr)
        return 2
    except (OSError, ValueError, csv.Error, LookupError) as exc:
        print(f"Conversion failed: {exc}", file=sys.stderr)
        return 2
    print(json.dumps(report, ensure_ascii=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
